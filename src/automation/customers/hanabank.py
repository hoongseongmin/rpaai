from pathlib import Path
from datetime import datetime
import re
import os
import time
import pandas as pd
import openpyxl
from openpyxl.utils import column_index_from_string

def process_hanabank(operation: str, customer: str, diff_file_path: list, output_dir: Path, options: dict) -> dict:
    if operation in ['hanabank_task1', 'doc_official', 'hanabank_task2_hybrid']:
        # 대시보드 UI에서 입력한 업체명 가져오기 (기본값: 에이텍)
        target_company = options.get('context_override', {}).get('추출_업체명', '에이텍')
        
        # AJAX로 업로드한 파일명 가져오기
        uploaded_filename = options.get('context_override', {}).get('원본_엑셀_업로드', '')
        
        # 대시보드 화면 렌더링 시 (파일 업로드 전 미리보기 요청 처리)
        if options.get('preview_only'):
            return {
                "success": True, 
                "output": "", 
                "summary": {
                    "message": "데이터 로드 완료. (엑셀 파일을 업로드하고 생성 버튼을 누르세요)", 
                    "context": {"원본_엑셀_업로드": uploaded_filename, "추출_업체명": target_company}
                }
            }
            
        # 💡 업로드된 파일 이름을 기반으로 실제 파일 경로 찾기
        file_path = None
        base_dir = output_dir.parent
        upload_dir = base_dir / "uploads" / "compare"
        
        if uploaded_filename:
            candidate_path = upload_dir / uploaded_filename
            if candidate_path.exists():
                file_path = candidate_path

        # 💡 원클릭 대시보드 등에서 파일명이 정상적으로 넘어오지 않았을 경우의 강력한 자동 매칭 폴백
        if not file_path:
            excel_files = list(upload_dir.glob("*.xlsx")) + list(upload_dir.glob("*.xls"))
            valid_files = [f for f in excel_files if "차이결과" not in f.name and not f.name.startswith("~")]
            if valid_files:
                file_path = max(valid_files, key=lambda p: p.stat().st_mtime)

        # 위에서 찾지 못했을 경우 기존 방식(diff_file_path) 폴백
        if not file_path:
            if diff_file_path:
                file_path = diff_file_path[0] if isinstance(diff_file_path, list) else diff_file_path
        
        if not file_path or not os.path.exists(file_path):
            return {"success": False, "message": "업로드된 엑셀 파일이 없습니다. [업로드] 버튼을 눌러 파일을 먼저 첨부해 주세요."}
            
        # 💡 방어 코드: 대시보드에서 파일 첨부를 안 해서 서버가 엉뚱한 과거 차이결과를 끌고 온 경우 즉시 차단
        if "차이결과" in file_path.name:
            return {"success": False, "message": "파일이 정상적으로 업로드되지 않았습니다. 반드시 파란색 [업로드] 버튼을 눌러 '✅ 업로드 완료!' 메시지를 확인한 뒤 생성해 주세요."}
            
        # 💡 핵심 로직: 구형 .xls 파일이 들어오면 서식을 보존한 채 신형 .xlsx로 백그라운드 자동 변환
        if file_path.suffix.lower() == '.xls':
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            excel = None
            wb_com = None
            try:
                excel = win32com.client.DispatchEx("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                wb_com = excel.Workbooks.Open(str(file_path.absolute()))
                
                # 💡 덮어쓰기 팝업창(먹통 현상) 원천 차단을 위해 매번 고유한 새 이름으로 변환 저장
                time_str_temp = datetime.now().strftime('%H%M%S_%f')
                new_path = file_path.with_name(f"{file_path.stem}_{time_str_temp}.xlsx")
                wb_com.SaveAs(str(new_path.absolute()), FileFormat=51) # 51은 xlsx 포맷 코드
                file_path = new_path
            except Exception as e:
                return {"success": False, "message": f".xls 엑셀 자동 변환 중 오류 발생: {str(e)}\n(서버 PC에 엑셀이 설치되어 있어야 합니다.)"}
            finally:
                if wb_com: wb_com.Close(False)
                if excel: excel.Quit()
                try: pythoncom.CoUninitialize()
                except: pass
        
        try:
            # 💡 [속도 측정용 타이머 시작]
            start_time = time.time()
            is_hybrid = (operation == 'hanabank_task2_hybrid')
            
            df_dict = None
            if is_hybrid:
                df_dict = pd.read_excel(file_path, sheet_name=None, header=None)
            
            # 1. 서식과 데이터(수식 등) 모두 보존하기 위해 data_only=False로 로드
            wb = openpyxl.load_workbook(file_path, data_only=False)
            
            if len(wb.worksheets) < 3:
                return {"success": False, "message": f"업로드한 엑셀 파일({file_path.name})에 최소 3개의 시트가 필요합니다."}
                
            ws1 = wb.worksheets[0]
            ws2 = wb.worksheets[1]
            ws3 = wb.worksheets[2]
            
            # 동적 시트명 추출 함수
            def get_dynamic_sheet_name(original_title):
                company_prefix = "ATEC" if target_company == "에이텍" else target_company
                match = re.search(r'(\(\d{2}월\)|_\d+)$', original_title)
                if match:
                    date_part = match.group(1)
                    name_part = original_title[:match.start()]
                    
                    if "휴일유지보수" in name_part and not date_part.startswith('_'):
                        new_title = f"({company_prefix})_{name_part}_{date_part}"
                    else:
                        new_title = f"({company_prefix})_{name_part}{date_part}"
                else:
                    new_title = f"({company_prefix})_{original_title}"
                    
                return new_title[:31]

            # 2. 기존 시트 뒤에 새 시트 복제 및 이름 변경 (서식 복사됨)
            new_ws1 = wb.copy_worksheet(ws1)
            new_ws1.title = get_dynamic_sheet_name(ws1.title)
            
            new_ws2 = wb.copy_worksheet(ws2)
            new_ws2.title = get_dynamic_sheet_name(ws2.title)
            
            new_ws3 = wb.copy_worksheet(ws3)
            new_ws3.title = get_dynamic_sheet_name(ws3.title)
            
            # 3. 불필요한 행 삭제 (역순 탐색)
            # 튜플 구성: (새 시트객체, 검색할 열 알파벳, 헤더보호를 위한 키워드)
            configs = [
                (new_ws1, 'G', '업체명'),
                (new_ws2, 'I', '업체'),
                (new_ws3, 'C', '업체명')
            ]
            
            if is_hybrid:
                hybrid_configs = [
                    (ws1, new_ws1, 'G', '업체명'),
                    (ws2, new_ws2, 'I', '업체'),
                    (ws3, new_ws3, 'C', '업체명')
                ]
                for orig_ws, new_ws, col_letter, header_kw in hybrid_configs:
                    sheet_name = orig_ws.title
                    if sheet_name not in df_dict:
                        continue
                        
                    df = df_dict[sheet_name]
                    col_idx = column_index_from_string(col_letter)
                    pandas_col_idx = col_idx - 1 # pandas는 0부터 시작
                    
                    if pandas_col_idx >= len(df.columns):
                        continue # 열이 비어있어 pandas가 무시한 경우 패스
                    
                    # 헤더 행 찾기 (1~20행 사이에서 키워드 검색, 타이틀 영역 보호)
                    header_row = 1
                    for idx in range(min(20, len(df))):
                        val = str(df.iloc[idx, pandas_col_idx]).strip() if pd.notna(df.iloc[idx, pandas_col_idx]) else ""
                        if header_kw in val:
                            header_row = idx + 1
                            break
                            
                    if header_row >= len(df):
                        continue
                    
                    # 지정된 열의 데이터만 Series로 추출 및 전처리
                    data_series = df.iloc[header_row:, pandas_col_idx].fillna('').astype(str)
                    
                    # 존재하는 마지막 데이터 행(actual_max) 찾기
                    non_empty_mask = data_series.str.strip() != ""
                    if not non_empty_mask.any():
                        actual_max = header_row
                    else:
                        last_idx = non_empty_mask[::-1].idxmax()
                        actual_max = last_idx + 1 # 엑셀 행은 인덱스 + 1
                            
                    # 데이터가 없는 하단 쓰레기 행들을 한 번에 광속 삭제
                    if new_ws.max_row > actual_max:
                        new_ws.delete_rows(actual_max + 1, new_ws.max_row - actual_max)
                            
                    # 💡 [Pandas 벡터 연산] 에이텍이 없는 행들을 0.001초 만에 식별하여 명단(Index) 도출
                    valid_series = data_series.loc[:actual_max - 1]
                    delete_mask = ~valid_series.str.contains(target_company)
                    rows_to_delete = sorted((valid_series[delete_mask].index + 1).tolist(), reverse=True)
                    
                    # 지워야 할 명단을 바탕으로 덩어리(블록) 생성
                    delete_blocks = []
                    if rows_to_delete:
                        block_start = rows_to_delete[0]
                        block_end = rows_to_delete[0]
                        for r in rows_to_delete[1:]:
                            if r == block_start - 1:
                                block_start = r
                            else:
                                delete_blocks.append((block_start, block_end - block_start + 1))
                                block_start = r
                                block_end = r
                        delete_blocks.append((block_start, block_end - block_start + 1))
                        
                    # 밑(하단)에서부터 위로 올라가며 덩어리(블록) 단위로 한 번에 광속 삭제
                    for start_row, amount in delete_blocks:
                        new_ws.delete_rows(start_row, amount)
                        
            else:
                for new_ws, col_letter, header_kw in configs:
                    col_idx = column_index_from_string(col_letter)
                            
                    # 헤더 행 찾기 (1~20행 사이에서 키워드 검색, 타이틀 영역 보호)
                    header_row = 1
                    for r in range(1, 21):
                        cell_val = str(new_ws.cell(row=r, column=col_idx).value).strip()
                        if header_kw in cell_val:
                            header_row = r
                            break
                            
                    # 💡 존재하는 셀만 가져와서 탐색하도록 완벽 최적화
                    actual_max = header_row
                    valid_rows = [r for (r, c) in new_ws._cells.keys() if c == col_idx and r > header_row]
                    valid_rows.sort(reverse=True)
                    
                    for r in valid_rows:
                        val = new_ws.cell(row=r, column=col_idx).value
                        if val is not None and str(val).strip() != "":
                            actual_max = r
                            break
                            
                    # 데이터가 없는 하단 쓰레기 행들을 한 번에 광속 삭제
                    if new_ws.max_row > actual_max:
                        new_ws.delete_rows(actual_max + 1, new_ws.max_row - actual_max)
                            
                    # 💡 [초고속 최적화] 지워야 할 행들을 '연속된 덩어리(블록)'로 묶어서 한 번에 삭제
                    delete_blocks = []
                    block_start = None
                    block_end = None
                    
                    for row in range(actual_max, header_row, -1):
                        cell_val = new_ws.cell(row=row, column=col_idx).value
                        val = str(cell_val).strip() if cell_val is not None else ""
                        
                        if target_company not in val:
                            if block_end is None:
                                block_end = row
                            block_start = row
                        else:
                            if block_end is not None:
                                delete_blocks.append((block_start, block_end - block_start + 1))
                                block_end = None
                                
                    if block_end is not None:
                        delete_blocks.append((block_start, block_end - block_start + 1))
                        
                    # 밑(하단)에서부터 위로 올라가며 덩어리(블록) 단위로 한 번에 광속 삭제
                    for start_row, amount in delete_blocks:
                        new_ws.delete_rows(start_row, amount)
                        
            # 4. 결과 파일 저장
            time_str = datetime.now().strftime('%H%M%S')
            output_name = f"02_(하이브리드)_하나은행_{time_str}.xlsx" if is_hybrid else f"02_(편집)_하나은행_{time_str}.xlsx"
            output_path = output_dir / output_name
            
            wb.save(output_path)
            
            elapsed_time = time.time() - start_time
            
            return {
                "success": True, 
                "output": output_name, 
                "summary": {"message": f"'{target_company}' 데이터만 추출하여 서식이 유지된 편집본이 생성되었습니다. ({'하이브리드 ' if is_hybrid else ''}처리시간: {elapsed_time:.2f}초)"}
            }
            
        except Exception as e:
            return {"success": False, "message": f"엑셀 처리 중 오류 발생: {str(e)}"}
            
    return {"success": False, "message": "지원하지 않는 작업입니다."}