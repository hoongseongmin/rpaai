import pandas as pd
from pathlib import Path
from datetime import datetime
import re
import json
from docxtpl import DocxTemplate

from ._common import parse_money, number_to_korean_currency

def export_custom_pdf(wb, file_path: Path, output_dir: Path):
    """우정사업본부 전용 커스텀 PDF 변환 (통합 엑셀 시트 분할 및 압축)"""
    import os
    import zipfile
    
    if "통합" in file_path.name and wb.Sheets.Count > 1:
        zip_filename = file_path.with_suffix('.zip').name
        zip_path = output_dir / zip_filename
        
        date_match = re.search(r'\((\d{4})년(\d{2})월분\)', file_path.name)
        target_year = date_match.group(1) if date_match else "YYYY"
        target_month = date_match.group(2) if date_match else "MM"
        short_year = target_year[-2:] if len(target_year) == 4 else "YY"
        
        custom_names = {
            1: f"공문_금융자동화기기(Ⅱ) 유지관리용역 이전설치 완료통보의 건({target_year}년{target_month}월분)",
            2: f"공문_금융자동화기기(Ⅱ) 유지관리용역 {short_year}.{target_month}월 이행실적 검사요청",
            3: f"청구공문_금융자동화기기(Ⅱ) 유지관리용역 이전실비({target_year}년{target_month}월분)",
            4: f"청구공문_금융자동화기기(Ⅱ) 유지관리용역 유지보수료({target_year}년{target_month}월분)"
        }

        pdf_files = []
        for i in range(1, wb.Sheets.Count + 1):
            sheet = wb.Sheets(i)
            base_pdf_name = custom_names.get(i, sheet.Name)
            safe_sheet_name = re.sub(r'[\\/*?:"<>|]', "", base_pdf_name)
            pdf_name = f"{safe_sheet_name}.pdf"
            pdf_temp_path = output_dir / pdf_name
            sheet.ExportAsFixedFormat(0, str(pdf_temp_path.absolute()))
            pdf_files.append(pdf_temp_path)
            
        with zipfile.ZipFile(str(zip_path.absolute()), 'w', zipfile.ZIP_DEFLATED) as zipf:
            for p_file in pdf_files: zipf.write(p_file, arcname=p_file.name)
                
        for p_file in pdf_files:
            try: os.remove(p_file)
            except: pass
            
        return {"success": True, "output": zip_filename, "summary": {"message": "4개 시트가 개별 PDF로 분할되어 압축(ZIP) 완료되었습니다."}}
    return None

def process_koreapost(operation: str, customer: str, diff_file_path: Path, output_dir: Path, options: dict) -> dict:
    """우정사업본부 전용 공문 및 청구서 생성 프로세스"""
    options = options or {}
    
    # === 🌟 통합 생성 (1~4번 엑셀 한 파일로 병합) ===
    if operation == 'kp_task_all':
        report_date = options.get('context_override', {}).get('청구연월', options.get('report_date', datetime.now().strftime('%Y-%m')))
        numbers = re.findall(r'\d+', report_date)
        now = datetime.now()
        target_year = int(numbers[0]) if len(numbers) > 0 else now.year
        target_month = int(numbers[1]) if len(numbers) > 1 else now.month
        
        merged_file_prefix = f"우정사업본부_통합_청구자료({target_year}년{target_month:02d}월분)"
        output_path = output_dir / f"{merged_file_prefix}.xlsx"
        
        if options.get('preview_only'):
            return {"success": True, "output": "", "summary": {"message": "1~4번 문서를 하나로 합쳐서 생성합니다. 하단의 [🚀 이 내용으로 문서 생성]을 클릭하세요.", "context": options.get('context_override', {})}}

        task_names = {
            'kp_task1': '1.이전설치공문',
            'kp_task2': '2.이행실적공문',
            'kp_task3': '3.이전실비청구서',
            'kp_task4': '4.유지보수료청구'
        }

        files_to_merge = []
        for task_op in ['kp_task1', 'kp_task2', 'kp_task3', 'kp_task4']:
            task_options = dict(options)
            task_options['preview_only'] = False
            task_ctx = dict(options.get('context_override', {}))
            
            # 각 작업에 맞게 분리된 키워드(3번용, 4번용)를 기본 키워드로 매핑
            if task_op == 'kp_task3':
                if '당초청구금액_3' in task_ctx: task_ctx['당초청구금액'] = task_ctx['당초청구금액_3']
                if '정산감액_3' in task_ctx: task_ctx['정산감액'] = task_ctx['정산감액_3']
                if '청구금액_3' in task_ctx: task_ctx['청구금액'] = task_ctx['청구금액_3']
            elif task_op == 'kp_task4':
                if '당초청구금액_4' in task_ctx: task_ctx['당초청구금액'] = task_ctx['당초청구금액_4']
                if '정산감액_4' in task_ctx: task_ctx['정산감액'] = task_ctx['정산감액_4']
                if '청구금액_4' in task_ctx: task_ctx['청구금액'] = task_ctx['청구금액_4']
                
            task_options['context_override'] = task_ctx
            
            sub_customer = f"koreapost_{task_op.replace('kp_task', '')}"
            res = process_koreapost(task_op, sub_customer, diff_file_path, output_dir, task_options)
            if res.get('success') and res.get('output'):
                files_to_merge.append((res['output'], task_names[task_op]))
            else:
                return {"success": False, "message": f"{task_op} 작업 실패: {res.get('message', '')}"}
        
        import win32com.client
        import pythoncom
        import os
        
        pythoncom.CoInitialize()
        excel = None
        wb_merged = None
        wb_sources = []
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            
            first_file, first_sheet_name = files_to_merge[0]
            wb_merged = excel.Workbooks.Open(os.path.abspath(first_file))
            
            # 첫 번째 파일의 시트 이름도 안전하게 변경
            sheet_idx = 1
            for i in range(1, wb_merged.Sheets.Count + 1):
                try: wb_merged.Sheets(i).Name = first_sheet_name if wb_merged.Sheets.Count == 1 else f"{first_sheet_name}_{sheet_idx}"
                except: pass
                sheet_idx += 1
            
            for file_path, sheet_name in files_to_merge[1:]:
                abs_path = os.path.abspath(file_path)
                wb_source = excel.Workbooks.Open(abs_path)
                wb_sources.append(wb_source)
                
                sheet_idx = 1
                for i in range(1, wb_source.Sheets.Count + 1):
                    source_sheet = wb_source.Sheets(i)
                    # 💡 핵심: 키워드 인수(After=) 대신 위치 인수로 전달해야 새 워크북으로 빠지지 않음
                    source_sheet.Copy(None, wb_merged.Sheets(wb_merged.Sheets.Count))
                    new_sheet = wb_merged.Sheets(wb_merged.Sheets.Count)
                    try:
                        new_sheet.Name = sheet_name if wb_source.Sheets.Count == 1 else f"{sheet_name}_{sheet_idx}"
                    except:
                        pass
                    sheet_idx += 1
                    
                wb_source.Close(False)
                wb_sources.remove(wb_source)
                try: os.remove(abs_path) # 병합 완료된 개별 파일은 삭제하여 폴더 정리
                except: pass
                    
            wb_merged.Sheets(1).Activate()
            wb_merged.SaveAs(os.path.abspath(output_path), FileFormat=51)
        except Exception as e:
            return {"success": False, "message": f"엑셀 병합 실패: {str(e)}"}
        finally:
            for wb in wb_sources:
                try: wb.Close(False)
                except: pass
            if wb_merged: wb_merged.Close(False)
            if excel: excel.Quit()
            try: pythoncom.CoUninitialize()
            except: pass
            
            try: os.remove(os.path.abspath(files_to_merge[0][0]))
            except: pass
            
        return {"success": True, "output": str(output_path), "summary": {"message": "1~4번이 모두 포함된 통합 엑셀 파일이 성공적으로 생성되었습니다."}}

    # === 1번: 이전설치 공문 (Excel) ===
    if operation == 'kp_task1':
        templates_dir = output_dir.parent / "templates"
        excel_settings_path = templates_dir / "우체국" / "[설정]_우체국_공문1_2.xlsx"
        
        try:
            df_settings = pd.read_excel(excel_settings_path, dtype=str)
            customer_row = df_settings[df_settings['고객사코드'] == customer]
            if customer_row.empty:
                base_customer = customer.split('_')[0]
                customer_row = df_settings[df_settings['고객사코드'] == base_customer]
                if customer_row.empty:
                    return {"success": False, "message": f"엑셀 설정 파일에 고객사코드 '{customer}' 항목이 존재하지 않습니다."}
            context = customer_row.fillna('').iloc[0].to_dict()
        except Exception as e:
            return {"success": False, "message": f"설정 파일 읽기 오류: {str(e)}"}
            
        if 'context_override' in options:
            context.update({k: v for k, v in options['context_override'].items() if str(v).strip()})

        report_date = context.get('청구연월', options.get('report_date', datetime.now().strftime('%Y-%m')))
        context['청구연월'] = report_date
        numbers = re.findall(r'\d+', report_date)
        now = datetime.now()
        target_year = int(numbers[0]) if len(numbers) > 0 else now.year
        target_month = int(numbers[1]) if len(numbers) > 1 else now.month
        
        context['청구연도'], context['청구월'] = str(target_year), f"{target_month:02d}월"
        context['전월연도'], context['전월'] = str(target_year if target_month > 1 else target_year - 1), f"{target_month - 1 if target_month > 1 else 12:02d}월"
        
        write_date_str = context.get('작성일자', '')
        if write_date_str:
            w_nums = re.findall(r'\d+', str(write_date_str))
            if len(w_nums) >= 3:
                w_year, w_month, w_day = int(w_nums[0]), int(w_nums[1]), int(w_nums[2])
            else:
                w_year, w_month, w_day = now.year, now.month, now.day
            try: target_write_date = datetime(w_year, w_month, w_day)
            except ValueError: target_write_date = now
        else:
            target_write_date = now

        context['작성일자'] = write_date_str if write_date_str else target_write_date.strftime('%Y년 %m월 %d일')
        context['작성일자_숫자'] = target_write_date.strftime('%Y%m%d')
        context['작성일자_점'] = target_write_date.strftime('%Y.%m.%d')
        context['작성연도_숫자'] = target_write_date.strftime('%Y')
        context['작성월_숫자'] = target_write_date.strftime('%m')
        context['작성일자_일'] = target_write_date.strftime('%d')

        kr_vars = {}
        for k, v in context.items():
            if any(mk in k for mk in ['금액', '단가', '가액', '세', '비용', '지급액']) and not k.endswith('_한글'):
                val_str = str(v).strip()
                if val_str:
                    amt = parse_money(val_str)
                    if amt != 0 or '0' in val_str:
                        context[k] = f"{amt:,}"
                kr_vars[f"{k}_한글"] = number_to_korean_currency(str(v))
        context.update(kr_vars)

        if options.get('preview_only'):
            replace_keys = list(context.keys())
            context['모든_변수_확인용'] = json.dumps({k: context.get(k, '') for k in replace_keys}, ensure_ascii=False, indent=2)
            return {"success": True, "output": "", "summary": {"message": "데이터 로드 완료", "context": context}}

        template_path = templates_dir / "우체국" / "[양식]_우체국_공문1.xlsx"
        if not template_path.exists():
            return {"success": False, "message": "엑셀 양식 파일을 찾을 수 없습니다."}
            
        full_month = f"{target_year}년{target_month:02d}월"
        file_prefix = f"공문_금융자동화기기(Ⅱ) 유지관리용역 이전설치 완료통보의 건({full_month}분)"
        output_path = output_dir / f"{file_prefix}.xlsx"
        
        try:
            import openpyxl
            wb = openpyxl.load_workbook(template_path, keep_links=False)
            pattern = re.compile(r'\{\{(.*?)\}\}')
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and pattern.findall(cell.value):
                            for match in pattern.findall(cell.value):
                                cell.value = cell.value.replace(f"{{{{{match.strip()}}}}}", str(context.get(match.strip(), '')))
            wb.save(output_path)
            return {"success": True, "output": str(output_path), "summary": {"message": "작업 완료: 이전설치 공문(Excel)이 생성되었습니다."}}
        except Exception as e:
            return {"success": False, "message": f"엑셀 템플릿 처리 중 오류 발생: {str(e)}"}

    # === 2번: 이행실적 공문 (Excel) ===
    if operation == 'kp_task2':
        templates_dir = output_dir.parent / "templates"
        excel_path = templates_dir / "우체국" / "[설정]_우체국_공문1_2.xlsx"
        
        if not excel_path.exists():
            return {"success": False, "message": f"다음 템플릿 파일을 찾지 못했습니다:\n\n- {excel_path.name}\n  (파이썬이 찾은 곳: {excel_path})"}
            
        try:
            df_settings = pd.read_excel(excel_path, dtype=str)
            customer_row = df_settings[df_settings['고객사코드'] == customer]
            
            if customer_row.empty:
                base_customer = customer.split('_')[0]
                customer_row = df_settings[df_settings['고객사코드'] == base_customer]
                if customer_row.empty:
                    return {"success": False, "message": f"엑셀 설정 파일에 고객사코드 '{customer}' 항목이 존재하지 않습니다."}
            
            context = customer_row.fillna('').iloc[0].to_dict()

            if not str(context.get('월유지비용', '')).strip():
                base_customer = customer.split('_')[0]
                fallback_row = df_settings[df_settings['고객사코드'] == base_customer]
                if not fallback_row.empty:
                    fallback_ctx = fallback_row.fillna('').iloc[0].to_dict()
                    if str(fallback_ctx.get('월유지비용', '')).strip():
                        context['월유지비용'] = fallback_ctx.get('월유지비용', '')
            
            for default_key in ['월유지비용', '미사용차감수량', '미사용차감금액', '실지급액']:
                if default_key not in context: context[default_key] = ''

            fixed_value_kp = context.get('월유지비용', '')

            if customer == 'koreapost_2':
                for default_key in ['장애_적기처리건수', '장애_지연처리건수', '만족도조사_건수', '만족도조사_총점', '만족도조사_평점', '정기점검_총대상수', '정기점검_완료수', '정기점검_달성률', '이전설치_총대상수', '이전설치_완료수', '이전설치_달성률']:
                    if default_key not in context: context[default_key] = ''

            template_name = '[양식]_우체국_공문2.xlsx'
            template_path = templates_dir / "우체국" / template_name
            
            if not template_path.exists() and not options.get('preview_only'):
                return {"success": False, "message": f"고객사 전용 엑셀 양식을 찾지 못했습니다:\n\n- {template_name}\n  (파이썬이 찾은 곳: {template_path})"}
                
            context['공문끝번호'] = options.get('doc_number', '')
            context['합계금액'] = ''
            context['공급가액'] = ''
            context['부가세'] = ''
            
            report_date = options.get('report_date', datetime.now().strftime('%Y-%m'))
            if 'context_override' in options and '청구연월' in options['context_override']:
                report_date = options['context_override']['청구연월']
                
            context['청구연월'] = report_date
            numbers = re.findall(r'\d+', report_date)
            now = datetime.now()
            target_year = int(numbers[0]) if len(numbers) > 0 else now.year
            target_month = int(numbers[1]) if len(numbers) > 1 else now.month
            
            context['청구연도'] = str(target_year)
            context['청구월'] = f"{target_month:02d}월"
            
            prev_month = target_month - 1 if target_month > 1 else 12
            prev_year = target_year if target_month > 1 else target_year - 1
            
            context['전월연도'] = str(prev_year)
            context['전월'] = f"{prev_month:02d}월"
            
            override = options.get('context_override', {})
            write_date_str = override.get('작성일자', '')
            if write_date_str:
                w_nums = re.findall(r'\d+', str(write_date_str))
                if len(w_nums) >= 3:
                    w_year, w_month, w_day = int(w_nums[0]), int(w_nums[1]), int(w_nums[2])
                else:
                    w_year, w_month, w_day = now.year, now.month, now.day
                try: target_write_date = datetime(w_year, w_month, w_day)
                except ValueError: target_write_date = now
            else:
                target_write_date = now

            context['작성일자'] = write_date_str if write_date_str else target_write_date.strftime('%Y년 %m월 %d일')
            context['작성일자_숫자'] = target_write_date.strftime('%Y%m%d')
            context['작성일자_점'] = target_write_date.strftime('%Y.%m.%d')
            context['작성연도_숫자'] = target_write_date.strftime('%Y')
            context['작성월_숫자'] = target_write_date.strftime('%m')
            context['작성일자_일'] = target_write_date.strftime('%d')
                
            if 'context_override' in options:
                cleaned_override = {k: v for k, v in options['context_override'].items() if str(v).strip()}
                context.update(cleaned_override)

            if fixed_value_kp is not None:
                context['월유지비용'] = fixed_value_kp

            base_amt = parse_money(context.get('월유지비용', ''))
            deduct_amt = parse_money(context.get('미사용차감금액', ''))
            final_amt = base_amt - deduct_amt
            
            context['월유지비용'] = f"{base_amt:,}"
            context['미사용차감금액'] = f"{deduct_amt:,}"
            context['실지급액'] = f"{final_amt:,}"

            kr_vars = {}
            for k, v in context.items():
                if any(mk in k for mk in ['금액', '단가', '가액', '세', '비용', '지급액']) and not k.endswith('_한글'):
                    val_str = str(v).strip()
                    if val_str:
                        amt = parse_money(val_str)
                        if amt != 0 or '0' in val_str:
                            context[k] = f"{amt:,}"
                    kr_vars[f"{k}_한글"] = number_to_korean_currency(str(v))
            context.update(kr_vars)
            
            if options.get('preview_only'):
                var_dict = {k: context.get(k, '') for k in context.keys()}
                context['모든_변수_확인용'] = json.dumps(var_dict, ensure_ascii=False, indent=2)
                return {"success": True, "output": "", "summary": {"message": "데이터 로드 완료", "context": context}}

            replace_keys = list(context.keys())
            for key, value in context.items():
                if isinstance(value, str):
                    for r_key in replace_keys:
                        if f"{{{r_key}}}" in value:
                            value = value.replace(f"{{{r_key}}}", str(context.get(r_key, '')))
                    context[key] = value

            body_list = []
            for key, value in context.items():
                if key.startswith('내용') and isinstance(value, str) and value.strip():
                    body_list.append(f"\t{value}" if '_' in key else value)
            context['본문리스트'] = body_list
            
            attach_keys = [k for k in context.keys() if k.startswith('붙임') and isinstance(context[k], str) and context[k].strip()]
            attach_keys.sort()
            attach_list = [context[k] for k in attach_keys]
                
            context['붙임리스트'] = attach_list
            context['붙임_첫번째'] = attach_list[0] if attach_list else ''
            context['붙임_나머지'] = attach_list[1:] if len(attach_list) > 1 else []
            
            attach_text = f"붙임 :\t{attach_list[0]}" if attach_list else ""
            for a in attach_list[1:]: attach_text += f"\n\t{a}"
            context['붙임_통합본'] = attach_text
            
            short_month = f"{str(target_year)[-2:]}.{target_month:02d}월"
            file_prefix = f"공문_금융자동화기기(Ⅱ) 유지관리용역 {short_month} 이행실적 검사요청"
            output_path = output_dir / f"{file_prefix}.xlsx"
            
            import openpyxl
            wb = openpyxl.load_workbook(template_path, keep_links=False)
            pattern = re.compile(r'\{\{(.*?)\}\}')
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and pattern.findall(cell.value):
                            for match in pattern.findall(cell.value):
                                cell.value = cell.value.replace(f"{{{{{match.strip()}}}}}", str(context.get(match.strip(), '')))
            wb.save(output_path)
            
            return {"success": True, "output": str(output_path), "summary": {"message": f"작업 완료: 우정사업본부 이행실적 공문(Excel)이 성공적으로 생성되었습니다.", "context": context}}
        except Exception as e:
            return {"success": False, "message": f"엑셀 공문 생성 중 오류 발생: {str(e)}"}

    # === 3, 4번: 이전실비 및 유지보수료 청구서 (Excel) ===
    if operation in ['kp_task3', 'kp_task4']:
        templates_dir = output_dir.parent / "templates"
        excel_settings_path = templates_dir / "우체국" / "[설정]_우체국_공문3_4.xlsx"
        
        context = {}
        try:
            df_kp = pd.read_excel(excel_settings_path, dtype=str)
            task_row = df_kp[df_kp['작업명'].str.contains('3|이전|kp_task3', na=False, case=False)] if operation == 'kp_task3' else df_kp[df_kp['작업명'].str.contains('4|유지보수|kp_task4', na=False, case=False)]
            if not task_row.empty: context = task_row.fillna('').iloc[0].to_dict()
        except Exception: pass
        
        for default_key in ['데이터건수', '특정금액', '합계금액', '공급가액', '부가세', '당초청구금액', '정산감액', '청구금액']:
            if default_key not in context: context[default_key] = ''
        
        report_date = options.get('context_override', {}).get('청구연월', options.get('report_date', datetime.now().strftime('%Y-%m')))
        context['청구연월'] = report_date
        numbers = re.findall(r'\d+', report_date)
        now = datetime.now()
        target_year = int(numbers[0]) if len(numbers) > 0 else now.year
        target_month = int(numbers[1]) if len(numbers) > 1 else now.month
        
        context['청구연도'], context['청구월'] = str(target_year), f"{target_month:02d}월"
        context['전월연도'], context['전월'] = str(target_year if target_month > 1 else target_year - 1), f"{target_month - 1 if target_month > 1 else 12:02d}월"
        
        write_date_str = options.get('context_override', {}).get('작성일자', '')
        if write_date_str:
            w_nums = re.findall(r'\d+', str(write_date_str))
            if len(w_nums) >= 3:
                w_year, w_month, w_day = int(w_nums[0]), int(w_nums[1]), int(w_nums[2])
            else:
                w_year, w_month, w_day = now.year, now.month, now.day
            try: target_write_date = datetime(w_year, w_month, w_day)
            except ValueError: target_write_date = now
        else:
            target_write_date = now

        context['작성일자'] = write_date_str if write_date_str else target_write_date.strftime('%Y년 %m월 %d일')
        context['작성일자_숫자'] = target_write_date.strftime('%Y%m%d')
        context['작성일자_점'] = target_write_date.strftime('%Y.%m.%d')
        context['작성연도_숫자'] = target_write_date.strftime('%Y')
        context['작성월_숫자'] = target_write_date.strftime('%m')
        context['작성일자_일'] = target_write_date.strftime('%d')
        
        if 'context_override' in options:
            cleaned_override = {k: v for k, v in options['context_override'].items() if str(v).strip()}
            context.update(cleaned_override)
            
        if operation == 'kp_task3':
            context['정산감액'] = '0'

        if '당초청구금액' in context or '정산감액' in context:
            base_amt = parse_money(context.get('당초청구금액', ''))
            deduct_amt = parse_money(context.get('정산감액', ''))
            context['당초청구금액'], context['정산감액'], context['청구금액'] = f"{base_amt:,}", f"{deduct_amt:,}", f"{base_amt - deduct_amt:,}"

        kr_vars = {}
        for k, v in context.items():
            if any(mk in k for mk in ['금액', '단가', '가액', '세', '감액', '비용', '지급액']) and not k.endswith('_한글'):
                val_str = str(v).strip()
                if val_str:
                    amt = parse_money(val_str)
                    if amt != 0 or '0' in val_str:
                        context[k] = f"{amt:,}"
                kr_vars[f"{k}_한글"] = number_to_korean_currency(str(v))
        context.update(kr_vars)
        
        replace_keys = list(context.keys())
        for key, value in context.items():
            if isinstance(value, str):
                for r_key in replace_keys:
                    if f"{{{{{r_key}}}}}" in value: value = value.replace(f"{{{{{r_key}}}}}", str(context.get(r_key, '')))
                context[key] = value

        df = pd.read_excel(diff_file_path) if diff_file_path else None
        context['표데이터'] = df.fillna('').to_dict('records') if df is not None else []

        if options.get('preview_only'):
            context['모든_변수_확인용'] = json.dumps({k: context.get(k, '') for k in replace_keys}, ensure_ascii=False, indent=2)
            return {"success": True, "output": "", "summary": {"message": "데이터 로드 완료", "context": context}}

        template_path = templates_dir / "우체국" / "[양식]_우체국_공문3_4.xlsx"
        if not template_path.exists(): return {"success": False, "message": "엑셀 양식 파일을 찾을 수 없습니다."}
            
        file_prefix = f"청구공문_금융자동화기기(Ⅱ) 유지관리용역 {'이전실비' if operation == 'kp_task3' else '유지보수료'}({target_year}년{target_month:02d}월분)"
        output_path = output_dir / f"{file_prefix}.xlsx"
        
        try:
            import openpyxl
            wb = openpyxl.load_workbook(template_path, keep_links=False)
            pattern = re.compile(r'\{\{(.*?)\}\}')
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and pattern.findall(cell.value):
                            for match in pattern.findall(cell.value): cell.value = cell.value.replace(f"{{{{{match.strip()}}}}}", str(context.get(match.strip(), '')))
            wb.save(output_path)
            return {"success": True, "output": str(output_path), "summary": {"message": "작업 완료: 우정사업본부 엑셀 청구서가 생성되었습니다."}}
        except Exception as e:
            return {"success": False, "message": f"엑셀 템플릿 처리 중 오류 발생: {str(e)}"}

    return {"success": False, "message": "해당 작업은 준비 중입니다."}